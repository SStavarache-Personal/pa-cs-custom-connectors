"""Builds the committed DOCX fixtures for the DocxTemplateGenerator connector.

The fixtures reproduce the XML conventions Microsoft Word writes (rsid
attributes, proofErr/bookmark markers splitting runs, w14:paraId, Word's exact
checkbox content-control markup, multi-section header/footer references) so the
engine is exercised against realistic run fragmentation. They are generated, not
saved by Word; see IMPLEMENTATION_NOTES.md for the remaining desktop-Word checks.

Usage (from the repository root):
    python3 connectors/DocxTemplateGenerator/fixtures/build_fixtures.py

Optional cross-check fixtures are produced when python-docx and LibreOffice
(soffice) are available; otherwise they are skipped and the committed copies stay.
"""

import json
import os
import shutil
import struct
import subprocess
import sys
import tempfile
import zipfile
import zlib
from xml.sax.saxutils import escape

HERE = os.path.dirname(os.path.abspath(__file__))
CONNECTOR = os.path.dirname(HERE)
EXAMPLES = os.path.join(CONNECTOR, "examples")
FIXED_TIME = (2024, 1, 1, 0, 0, 0)

NS = (
    'xmlns:wpc="http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:o="urn:schemas-microsoft-com:office:office" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
    'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
    'xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:wp14="http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:w10="urn:schemas-microsoft-com:office:word" '
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" '
    'xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml" '
    'xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup" '
    'xmlns:wpi="http://schemas.microsoft.com/office/word/2010/wordprocessingInk" '
    'xmlns:wne="http://schemas.microsoft.com/office/word/2006/wordml" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'mc:Ignorable="w14 w15 wp14"'
)
DECL = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'

_para_counter = [0x1A2B0000]


def para_id():
    _para_counter[0] += 1
    return "%08X" % _para_counter[0]


# ----------------------------------------------------------------------------
# WordprocessingML helpers
# ----------------------------------------------------------------------------

def r(text, rpr="", rsid="00B1C2D3"):
    props = "<w:rPr>%s</w:rPr>" % rpr if rpr else ""
    space = ' xml:space="preserve"' if text != text.strip() else ""
    return '<w:r w:rsidR="%s">%s<w:t%s>%s</w:t></w:r>' % (rsid, props, space, escape(text))


def proof(kind="spellStart"):
    return '<w:proofErr w:type="%s"/>' % kind


def p(*content, ppr="", rsid="00A4F1E2"):
    props = "<w:pPr>%s</w:pPr>" % ppr if ppr else ""
    return '<w:p w14:paraId="%s" w14:textId="77777777" w:rsidR="%s" w:rsidRDefault="%s">%s%s</w:p>' % (
        para_id(), rsid, rsid, props, "".join(content))


def checkbox(tag, sdt_id, checked=False, rpr_extra=""):
    glyph = "☒" if checked else "☐"
    return (
        "<w:sdt><w:sdtPr>%s<w:tag w:val=\"%s\"/><w:id w:val=\"%d\"/>"
        "<w14:checkbox><w14:checked w14:val=\"%d\"/>"
        "<w14:checkedState w14:val=\"2612\" w14:font=\"MS Gothic\"/>"
        "<w14:uncheckedState w14:val=\"2610\" w14:font=\"MS Gothic\"/>"
        "</w14:checkbox></w:sdtPr><w:sdtEndPr/><w:sdtContent>"
        "<w:r w:rsidR=\"00C7D8E9\"><w:rPr><w:rFonts w:ascii=\"MS Gothic\" w:eastAsia=\"MS Gothic\" "
        "w:hAnsi=\"MS Gothic\" w:hint=\"eastAsia\"/>%s</w:rPr><w:t>%s</w:t></w:r>"
        "</w:sdtContent></w:sdt>"
    ) % ("<w:rPr>%s</w:rPr>" % rpr_extra if rpr_extra else "", tag, sdt_id, 1 if checked else 0, rpr_extra, glyph)


def tc(width, *paragraphs, tcpr_extra=""):
    return '<w:tc><w:tcPr><w:tcW w:w="%d" w:type="dxa"/>%s</w:tcPr>%s</w:tc>' % (width, tcpr_extra, "".join(paragraphs))


def tr(*cells, trpr=""):
    props = "<w:trPr>%s</w:trPr>" % trpr if trpr else ""
    return '<w:tr w:rsidR="00D1E2F3" w14:paraId="%s" w14:textId="77777777" w:rsidTr="00D1E2F3">%s%s</w:tr>' % (
        para_id(), props, "".join(cells))


def tbl(widths, *rows, style="ProductTable", look="04A0"):
    grid = "".join('<w:gridCol w:w="%d"/>' % w for w in widths)
    return (
        '<w:tbl><w:tblPr><w:tblStyle w:val="%s"/><w:tblW w:w="%d" w:type="dxa"/>'
        '<w:tblLook w:val="%s" w:firstRow="1" w:lastRow="1" w:firstColumn="0" w:lastColumn="0" w:noHBand="0" w:noVBand="1"/>'
        '</w:tblPr><w:tblGrid>%s</w:tblGrid>%s</w:tbl>'
    ) % (style, sum(widths), look, grid, "".join(rows))


def field_page():
    return (
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        '<w:r><w:t>1</w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
    )


def inline_image(rel_id, doc_pr_id, cx=914400, cy=457200):
    return (
        '<w:r><w:rPr><w:noProof/></w:rPr><w:drawing>'
        '<wp:inline distT="0" distB="0" distL="0" distR="0">'
        '<wp:extent cx="%d" cy="%d"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
        '<wp:docPr id="%d" name="Picture %d" descr="Static logo"/>'
        '<wp:cNvGraphicFramePr><a:graphicFrameLocks xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" noChangeAspect="1"/></wp:cNvGraphicFramePr>'
        '<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        '<pic:pic xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        '<pic:nvPicPr><pic:cNvPr id="0" name="logo.png"/><pic:cNvPicPr/></pic:nvPicPr>'
        '<pic:blipFill><a:blip r:embed="%s"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
        '<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="%d" cy="%d"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr>'
        '</pic:pic></a:graphicData></a:graphic></wp:inline></w:drawing></w:r>'
    ) % (cx, cy, doc_pr_id, doc_pr_id, rel_id, cx, cy)


def sect_pr(refs, landscape=False, title_page=False):
    references = "".join(
        '<w:%sReference w:type="%s" r:id="%s"/>' % (kind, variant, rid) for kind, variant, rid in refs)
    size = '<w:pgSz w:w="15840" w:h="12240" w:orient="landscape"/>' if landscape else '<w:pgSz w:w="12240" w:h="15840"/>'
    return (
        '<w:sectPr w:rsidR="00A4F1E2">%s%s'
        '<w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440" w:header="708" w:footer="708" w:gutter="0"/>'
        '<w:cols w:space="708"/>%s<w:docGrid w:linePitch="360"/></w:sectPr>'
    ) % (references, size, "<w:titlePg/>" if title_page else "")


def png_bytes(width=16, height=8, rgb=(0x2B, 0x57, 0x9A)):
    raw = b"".join(b"\x00" + bytes(rgb) * width for _ in range(height))
    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw, 9)) + chunk(b"IEND", b""))


# ----------------------------------------------------------------------------
# Package parts
# ----------------------------------------------------------------------------

STYLES = DECL + (
    '<w:styles xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" mc:Ignorable="w14">'
    '<w:docDefaults><w:rPrDefault><w:rPr><w:rFonts w:ascii="Calibri" w:eastAsia="Calibri" w:hAnsi="Calibri" w:cs="Calibri"/>'
    '<w:sz w:val="22"/><w:szCs w:val="22"/><w:lang w:val="en-CA" w:eastAsia="en-US" w:bidi="ar-SA"/></w:rPr></w:rPrDefault>'
    '<w:pPrDefault><w:pPr><w:spacing w:after="160" w:line="259" w:lineRule="auto"/></w:pPr></w:pPrDefault></w:docDefaults>'
    '<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/><w:qFormat/></w:style>'
    '<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/>'
    '<w:uiPriority w:val="9"/><w:qFormat/><w:pPr><w:keepNext/><w:spacing w:before="240" w:after="120"/><w:outlineLvl w:val="0"/></w:pPr>'
    '<w:rPr><w:color w:val="2B579A"/><w:sz w:val="32"/></w:rPr></w:style>'
    '<w:style w:type="paragraph" w:customStyle="1" w:styleId="Disclaimer"><w:name w:val="Disclaimer"/><w:basedOn w:val="Normal"/>'
    '<w:qFormat/><w:pPr><w:pBdr><w:left w:val="single" w:sz="12" w:space="4" w:color="C00000"/></w:pBdr><w:spacing w:before="120" w:after="120"/></w:pPr>'
    '<w:rPr><w:i/><w:color w:val="595959"/><w:sz w:val="18"/></w:rPr></w:style>'
    '<w:style w:type="paragraph" w:styleId="ListParagraph"><w:name w:val="List Paragraph"/><w:basedOn w:val="Normal"/>'
    '<w:uiPriority w:val="34"/><w:qFormat/><w:pPr><w:ind w:left="720"/><w:contextualSpacing/></w:pPr></w:style>'
    '<w:style w:type="character" w:default="1" w:styleId="DefaultParagraphFont"><w:name w:val="Default Paragraph Font"/><w:uiPriority w:val="1"/><w:semiHidden/></w:style>'
    '<w:style w:type="character" w:customStyle="1" w:styleId="AccentEmphasis"><w:name w:val="Accent Emphasis"/>'
    '<w:basedOn w:val="DefaultParagraphFont"/><w:qFormat/><w:rPr><w:b/><w:color w:val="C55A11"/></w:rPr></w:style>'
    '<w:style w:type="table" w:default="1" w:styleId="TableNormal"><w:name w:val="Normal Table"/><w:semiHidden/>'
    '<w:tblPr><w:tblInd w:w="0" w:type="dxa"/><w:tblCellMar><w:top w:w="0" w:type="dxa"/><w:left w:w="108" w:type="dxa"/>'
    '<w:bottom w:w="0" w:type="dxa"/><w:right w:w="108" w:type="dxa"/></w:tblCellMar></w:tblPr></w:style>'
    '<w:style w:type="table" w:customStyle="1" w:styleId="ProductTable"><w:name w:val="Product Table"/><w:basedOn w:val="TableNormal"/>'
    '<w:tblPr><w:tblBorders><w:top w:val="single" w:sz="4" w:space="0" w:color="8EAADB"/><w:left w:val="single" w:sz="4" w:space="0" w:color="8EAADB"/>'
    '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="8EAADB"/><w:right w:val="single" w:sz="4" w:space="0" w:color="8EAADB"/>'
    '<w:insideH w:val="single" w:sz="4" w:space="0" w:color="8EAADB"/><w:insideV w:val="single" w:sz="4" w:space="0" w:color="8EAADB"/></w:tblBorders></w:tblPr>'
    '<w:tblStylePr w:type="firstRow"><w:rPr><w:b/><w:color w:val="FFFFFF"/></w:rPr><w:tcPr><w:shd w:val="clear" w:color="auto" w:fill="2B579A"/></w:tcPr></w:tblStylePr>'
    '<w:tblStylePr w:type="lastRow"><w:rPr><w:b/></w:rPr><w:tcPr><w:tcBorders><w:top w:val="double" w:sz="4" w:space="0" w:color="2B579A"/></w:tcBorders></w:tcPr></w:tblStylePr>'
    '</w:style>'
    '<w:style w:type="numbering" w:default="1" w:styleId="NoList"><w:name w:val="No List"/><w:semiHidden/></w:style>'
    '</w:styles>'
)

NUMBERING = DECL + (
    '<w:numbering xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    '<w:abstractNum w:abstractNumId="0"><w:multiLevelType w:val="hybridMultilevel"/>'
    '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="•"/><w:lvlJc w:val="left"/>'
    '<w:pPr><w:ind w:left="720" w:hanging="360"/></w:pPr><w:rPr><w:rFonts w:ascii="Symbol" w:hAnsi="Symbol" w:hint="default"/></w:rPr></w:lvl>'
    '</w:abstractNum><w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num></w:numbering>'
)

SETTINGS = DECL + (
    '<w:settings xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" '
    'xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml" mc:Ignorable="w14 w15">'
    '<w:zoom w:percent="100"/>%s<w:proofState w:spelling="clean" w:grammar="clean"/><w:defaultTabStop w:val="720"/>'
    '<w:characterSpacingControl w:val="doNotCompress"/>'
    '<w:compat><w:compatSetting w:name="compatibilityMode" w:uri="http://schemas.microsoft.com/office/word" w:val="15"/></w:compat>'
    '<w:rsids><w:rsidRoot w:val="00A4F1E2"/><w:rsid w:val="00A4F1E2"/><w:rsid w:val="00B1C2D3"/><w:rsid w:val="00C7D8E9"/><w:rsid w:val="00D1E2F3"/></w:rsids>'
    '<w:themeFontLang w:val="en-CA"/><w:decimalSymbol w:val="."/><w:listSeparator w:val=","/>'
    '<w15:docId w15:val="{6A1F3C1E-3B6F-4C55-9E64-1E2C7AA1D001}"/></w:settings>'
)

FONTS = DECL + (
    '<w:fonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    '<w:font w:name="Calibri"><w:panose1 w:val="020F0502020204030204"/><w:charset w:val="00"/><w:family w:val="swiss"/><w:pitch w:val="variable"/></w:font>'
    '<w:font w:name="MS Gothic"><w:altName w:val="ＭＳ ゴシック"/><w:charset w:val="80"/><w:family w:val="modern"/><w:pitch w:val="fixed"/></w:font>'
    '<w:font w:name="Symbol"><w:charset w:val="02"/><w:family w:val="roman"/><w:pitch w:val="variable"/></w:font>'
    '</w:fonts>'
)

CORE = DECL + (
    '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
    'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
    'xmlns:dcmitype="http://purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
    '<dc:title>Synthetic template fixture</dc:title><dc:creator>DocxTemplateGenerator fixtures</dc:creator>'
    '<dcterms:created xsi:type="dcterms:W3CDTF">2024-01-01T00:00:00Z</dcterms:created>'
    '<dcterms:modified xsi:type="dcterms:W3CDTF">2024-01-01T00:00:00Z</dcterms:modified></cp:coreProperties>'
)

APP = DECL + (
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
    'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
    '<Template>Normal.dotm</Template><Application>Microsoft Office Word</Application><DocSecurity>0</DocSecurity>'
    '<AppVersion>16.0000</AppVersion></Properties>'
)


def story(root, body):
    return DECL + "<w:%s %s>%s</w:%s>" % (root, NS, body, root)


def build_package(path, document_body, headers=None, footers=None, even_and_odd=False, include_image=False):
    """Writes a .docx. headers/footers map part file name -> inner XML."""
    headers = headers or {}
    footers = footers or {}
    overrides = [
        ("/word/document.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"),
        ("/word/styles.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"),
        ("/word/settings.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.settings+xml"),
        ("/word/fontTable.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.fontTable+xml"),
        ("/word/numbering.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml"),
        ("/docProps/core.xml", "application/vnd.openxmlformats-package.core-properties+xml"),
        ("/docProps/app.xml", "application/vnd.openxmlformats-officedocument.extended-properties+xml"),
    ]
    rels = [
        ("rId1", "styles", "styles.xml"),
        ("rId2", "settings", "settings.xml"),
        ("rId3", "fontTable", "fontTable.xml"),
        ("rId4", "numbering", "numbering.xml"),
    ]
    for name in headers:
        overrides.append(("/word/" + name, "application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml"))
    for name in footers:
        overrides.append(("/word/" + name, "application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml"))
    content_types = DECL + (
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="png" ContentType="image/png"/>'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>%s</Types>'
    ) % "".join('<Override PartName="%s" ContentType="%s"/>' % o for o in overrides)
    base = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
    rel_xml = "".join('<Relationship Id="%s" Type="%s%s" Target="%s"/>' % (i, base, t, target) for i, t, target in rels)
    doc_rels = (DECL + '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">%s%s%s%s</Relationships>') % (
        rel_xml,
        "".join('<Relationship Id="rIdH%s" Type="%sheader" Target="%s"/>' % (n[len("header"):-4], base, n) for n in headers),
        "".join('<Relationship Id="rIdF%s" Type="%sfooter" Target="%s"/>' % (n[len("footer"):-4], base, n) for n in footers),
        '<Relationship Id="rIdImg1" Type="%simage" Target="media/image1.png"/>' % base if include_image else "",
    )
    root_rels = DECL + (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties" Target="docProps/app.xml"/>'
        '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
        '</Relationships>'
    )
    entries = [
        ("[Content_Types].xml", content_types),
        ("_rels/.rels", root_rels),
        ("word/document.xml", story("document", "<w:body>%s</w:body>" % document_body)),
        ("word/_rels/document.xml.rels", doc_rels),
    ]
    for name, inner in headers.items():
        entries.append(("word/" + name, story("hdr", inner)))
    for name, inner in footers.items():
        entries.append(("word/" + name, story("ftr", inner)))
    entries += [
        ("word/styles.xml", STYLES),
        ("word/settings.xml", SETTINGS % ("<w:evenAndOddHeaders/>" if even_and_odd else "")),
        ("word/fontTable.xml", FONTS),
        ("word/numbering.xml", NUMBERING),
        ("docProps/core.xml", CORE),
        ("docProps/app.xml", APP),
    ]
    if include_image:
        entries.append(("word/media/image1.png", png_bytes()))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries:
            info = zipfile.ZipInfo(name, FIXED_TIME)
            info.compress_type = zipfile.ZIP_STORED if name.endswith(".png") else zipfile.ZIP_DEFLATED
            archive.writestr(info, data.encode("utf-8") if isinstance(data, str) else data)


# ----------------------------------------------------------------------------
# Example template (also the user-guide sample)
# ----------------------------------------------------------------------------

def example_template(path):
    bold = "<w:b/>"
    body = "".join([
        # Title: {CompanyName} fragmented across three runs with a proofing marker.
        p(r("Product Summary — "), r("{Comp", rsid="00B1C2D4"), proof(), r("any", rsid="00B1C2D5"), proof("spellEnd"),
          r("Name}", rsid="00B1C2D6"), ppr='<w:pStyle w:val="Heading1"/>'),
        # Mixed formatting around and inside tags.
        p(r("Prepared for "), r("{", bold), r("Customer.Name", bold + "<w:i/>"), r("}", bold), r(" ("),
          r("{Customer.City}", '<w:rStyle w:val="AccentEmphasis"/>'), r(") on {ReportDate}.")),
        p(r("Approved: "), checkbox("IsApproved", -1758194170), r("    Requires review: "), checkbox("RequiresReview", 1452821366)),
        tbl([3400, 1900, 1700, 1600],
            tr(tc(3400, p(r("Product"))), tc(1900, p(r("DIN"))), tc(1700, p(r("Price"))), tc(1600, p(r("Covered"))),
               trpr='<w:tblHeader/>'),
            tr(tc(3400, p(r("{#Products}"), r("{Name}"))),
               tc(1900, p(r("{DIN}"))),
               tc(1700, p(r("{Price"), r("Formatted}"), ppr='<w:jc w:val="right"/>')),
               tc(1600, p(checkbox("IsCovered", 2087655432), r("{/Products}"), ppr='<w:jc w:val="center"/>'))),
            tr(tc(3400, p(r("{#ShowDiscountRow}Discount"))), tc(1900, p()),
               tc(1700, p(r("{DiscountFormatted}"), ppr='<w:jc w:val="right"/>')), tc(1600, p(r("{/ShowDiscountRow}")))),
            tr(tc(3400, p(r("Total"))), tc(1900, p()), tc(1700, p(r("{TotalFormatted}"), ppr='<w:jc w:val="right"/>')), tc(1600, p())),
        ),
        p(),
        p(r("{#ShowDisclaimer}")),
        p(r("This summary was prepared for {CompanyName}; prices are indicative and exclude taxes."), ppr='<w:pStyle w:val="Disclaimer"/>'),
        p(r("Coverage is subject to the plan rules."), ppr='<w:pStyle w:val="ListParagraph"/><w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr>'),
        p(r("Contact your representative for details."), ppr='<w:pStyle w:val="ListParagraph"/><w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr>'),
        p(r("{/ShowDisclaimer}")),
        p(r("{#ShowInternalNotes}")),
        p(r("Internal: margin review pending.", "<w:color w:val=\"C00000\"/>")),
        p(r("{/ShowInternalNotes}")),
        p(inline_image("rIdImg1", 1), ppr='<w:jc w:val="center"/>'),
        p(r("Notes: ", bold), '<w:bookmarkStart w:id="0" w:name="_GoBack"/>', r("{Notes}"), '<w:bookmarkEnd w:id="0"/>'),
        p(ppr=sect_pr([("header", "first", "rIdH1"), ("header", "default", "rIdH2"), ("header", "even", "rIdH3"),
                       ("footer", "default", "rIdF1"), ("footer", "first", "rIdF2")], title_page=True)),
        p(r("Appendix"), ppr='<w:pStyle w:val="Heading1"/>'),
        p(r("Prices in this appendix use the same data as the summary for {CompanyName}")),
        sect_pr([("header", "default", "rIdH4")], landscape=True),
    ])
    headers = {
        "header1.xml": p(r("CONFIDENTIAL — {CompanyName}", "<w:b/><w:color w:val=\"C00000\"/>")),
        "header2.xml": p(r("{CompanyName}"), r(" | Product Summary | Approved "), checkbox("IsApproved", 511223344)),
        "header3.xml": p(r("Even page — {Customer.Name}")),
        "header4.xml": p(r("Appendix — {ReportDate}"), ppr='<w:jc w:val="right"/>'),
    }
    footers = {
        "footer1.xml": tbl([4600, 4600],
                           tr(tc(4600, p(r("Contact"))), tc(4600, p(r("Email"))), trpr='<w:tblHeader/>'),
                           tr(tc(4600, p(r("{#Contacts}{Name}"))), tc(4600, p(r("{Email}{/Contacts}")))),
                           style="ProductTable") + p(r("Page "), field_page(), ppr='<w:jc w:val="center"/>'),
        "footer2.xml": p(r("{FooterNote}", "<w:sz w:val=\"16\"/>")),
    }
    build_package(path, body, headers, footers, even_and_odd=True, include_image=True)


EXAMPLE_DATA = {
    "CompanyName": "Exemple Pharma Inc.",
    "Customer": {"Name": "Hôpital Sainte-Thérèse", "City": "Montréal"},
    "ReportDate": "2026-10-11",
    "IsApproved": True,
    "RequiresReview": False,
    "Products": [
        {"Name": "Atorvastatin", "DIN": "02456789", "PriceFormatted": "$12.50", "IsCovered": True},
        {"Name": "Metformin", "DIN": "02345678", "PriceFormatted": "$8.25", "IsCovered": False},
        {"Name": "Café & Crème <Test> \"quoted\" l'été", "DIN": "00000001", "PriceFormatted": "12,50 $", "IsCovered": True},
    ],
    "ShowDiscountRow": False,
    "DiscountFormatted": "-$1.00",
    "TotalFormatted": "$33.25",
    "ShowDisclaimer": True,
    "ShowInternalNotes": False,
    "Notes": "First line of the notes.\nSecond line — Unicode ✓ 東京 \U0001F680",
    "Contacts": [{"Name": "A. Tremblay", "Email": "a.tremblay@example.com"}],
    "FooterNote": "Generated by Power Automate",
}


# ----------------------------------------------------------------------------
# Independent cross-check fixtures
# ----------------------------------------------------------------------------

def python_docx_fixture(path):
    try:
        import docx
    except ImportError:
        print("python-docx not installed; skipping", path)
        return False
    document = docx.Document()
    document.add_heading("Order {OrderNumber}", level=1)
    paragraph = document.add_paragraph("Customer: ")
    run = paragraph.add_run("{Customer")
    run.bold = True
    paragraph.add_run(".Name}")
    document.add_paragraph("{#ShowTerms}")
    document.add_paragraph("Terms apply to {Customer.Name}.")
    document.add_paragraph("{/ShowTerms}")
    table = document.add_table(rows=2, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ["Item", "Qty", "Amount"]):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, ["{#Lines}{Item}", "{Qty}", "{Amount}{/Lines}"]):
        cell.text = text
    section = document.sections[0]
    section.header.paragraphs[0].text = "Header for {OrderNumber}"
    section.footer.paragraphs[0].text = "Footer {Customer.Name}"
    import datetime
    document.core_properties.created = datetime.datetime(2024, 1, 1)
    document.core_properties.modified = datetime.datetime(2024, 1, 1)
    document.save(path)
    return True


INDEPENDENT_DATA = {
    "OrderNumber": "SO-1001",
    "Customer": {"Name": "Société Générale Ltée"},
    "ShowTerms": True,
    "Lines": [{"Item": "Widget", "Qty": 2, "Amount": "$4.00"}, {"Item": "Gadget", "Qty": 1, "Amount": "$9.99"}],
}


def libreoffice_fixture(source, path):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("LibreOffice not installed; skipping", path)
        return False
    work = tempfile.mkdtemp()
    try:
        copy = os.path.join(work, "input.docx")
        shutil.copy(source, copy)
        out = os.path.join(work, "out")
        os.mkdir(out)
        subprocess.run([soffice, "--headless", "-env:UserInstallation=file://" + os.path.join(work, "profile"),
                        "--convert-to", "docx:MS Word 2007 XML", "--outdir", out, copy],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
        shutil.copy(os.path.join(out, "input.docx"), path)
        return True
    finally:
        shutil.rmtree(work, ignore_errors=True)


def write_json(path, data):
    with open(path, "w", encoding="utf-8", newline="\n") as handle:
        json.dump(data, handle, ensure_ascii=False, indent=2)
        handle.write("\n")


def write_scenarios(example_path):
    """Seeds the shared harness/workbench scenarios (connectors/<Name>/tests/*.json)."""
    import base64
    tests = os.path.join(CONNECTOR, "tests")
    os.makedirs(tests, exist_ok=True)
    template = base64.b64encode(open(example_path, "rb").read()).decode("ascii")
    data = json.dumps(EXAMPLE_DATA, ensure_ascii=False, separators=(",", ":"))
    missing = dict(EXAMPLE_DATA)
    del missing["CompanyName"]
    scenarios = {
        "generate-example.json": {
            "name": "GenerateDocument fills the example template",
            "operationId": "GenerateDocument", "relativePath": "/generate-document",
            "requestBody": {"templateBase64": template, "dataJson": data, "fileName": "ProductSummary.docx", "strictMode": True},
            "expected": {"statusCode": 200, "outboundRequestCount": 0, "bodyJson": {
                "fileName": "ProductSummary.docx",
                "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                "warnings": [],
                "statistics": {"textFieldsReplaced": 24, "tableRowsCreated": 4, "conditionalBlocksEvaluated": 3, "checkboxesUpdated": 6, "storyPartsProcessed": 7}}},
        },
        "validate-example-with-data.json": {
            "name": "ValidateTemplate checks structure and data without producing a file",
            "operationId": "ValidateTemplate", "relativePath": "/validate-template",
            "requestBody": {"templateBase64": template, "dataJson": data},
            "expected": {"statusCode": 200, "outboundRequestCount": 0, "bodyJson": {"valid": True, "dataValidated": True, "errors": []}},
        },
        "list-placeholders-example.json": {
            "name": "ListPlaceholders inventories the example template",
            "operationId": "ListPlaceholders", "relativePath": "/list-placeholders",
            "requestBody": {"templateBase64": template},
            "expected": {"statusCode": 200, "outboundRequestCount": 0, "bodyJson": {"errors": []},
                         "jsonPathEquals": {"fields[0].dataPath": "CompanyName", "fields[0].type": "text"}},
        },
        "missing-field-strict.json": {
            "name": "Strict mode rejects a missing value with its location",
            "operationId": "GenerateDocument", "relativePath": "/generate-document",
            "requestBody": {"templateBase64": template, "dataJson": json.dumps(missing, ensure_ascii=False, separators=(",", ":"))},
            "expected": {"statusCode": 422, "outboundRequestCount": 0, "bodyJson": {"error": {
                "code": "MISSING_FIELD", "field": "CompanyName", "part": "word/document.xml", "location": "paragraph 1"}}},
        },
        "invalid-base64.json": {
            "name": "Invalid base64 is a 400 error",
            "operationId": "GenerateDocument", "relativePath": "/generate-document",
            "requestBody": {"templateBase64": "not base64!", "dataJson": "{}"},
            "expected": {"statusCode": 400, "outboundRequestCount": 0, "bodyJson": {"error": {"code": "INVALID_BASE64"}}},
        },
    }
    for name, scenario in scenarios.items():
        scenario = dict(scenario)
        scenario = {"name": scenario.pop("name"), "includeInAutomatedRun": True, "operationId": scenario.pop("operationId"),
                    "method": "POST", **scenario}
        write_json(os.path.join(tests, name), scenario)


def main():
    os.makedirs(EXAMPLES, exist_ok=True)
    example = os.path.join(EXAMPLES, "ProductSummaryTemplate.docx")
    example_template(example)
    write_json(os.path.join(EXAMPLES, "ProductSummary.data.json"), EXAMPLE_DATA)
    print("wrote", example)
    write_scenarios(example)

    independent = os.path.join(HERE, "independent-python-docx.docx")
    if python_docx_fixture(independent):
        write_json(os.path.join(HERE, "independent.data.json"), INDEPENDENT_DATA)
        print("wrote", independent)
        resaved = os.path.join(HERE, "independent-libreoffice.docx")
        if libreoffice_fixture(independent, resaved):
            print("wrote", resaved)
    return 0


if __name__ == "__main__":
    sys.exit(main())
